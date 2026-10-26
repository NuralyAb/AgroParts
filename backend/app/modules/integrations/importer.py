"""Массовая загрузка ассортимента: Excel/CSV от магазина и пакеты синхронизации из 1С.

Каждая строка обрабатывается в собственном SAVEPOINT: ошибка в строке не ломает весь файл.
Остаток устанавливается через inventory.set_count — с записью в журнал и проверкой резервов.
"""

import csv
import io
from dataclasses import dataclass, field

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.errors import DomainError
from app.modules.catalog.normalize import normalize_number
from app.modules.integrations.matching import auto_match
from app.modules.inventory import service as inv
from app.modules.inventory.models import Offer
from app.modules.shops.models import Shop, Warehouse

COLUMN_ALIASES = {
    "sku": {"sku", "код", "код товара", "внутренний код", "артикул магазина"},
    "article": {"article", "артикул", "oem", "каталожный номер", "номер", "номер детали", "part number"},
    "title": {"title", "name", "наименование", "название", "товар"},
    "manufacturer": {"manufacturer", "brand", "производитель", "бренд"},
    "price": {"price", "цена", "цена, тг", "цена тг", "розница"},
    "qty": {"qty", "quantity", "количество", "кол-во", "остаток", "наличие"},
    "bin": {"bin", "ячейка", "место", "полка"},
    "warehouse": {"warehouse", "склад"},
}
TEMPLATE_HEADER = ["Артикул", "Наименование", "Производитель", "Цена", "Количество", "Ячейка", "Код"]


@dataclass
class ImportReport:
    rows: int = 0
    created: int = 0
    updated: int = 0
    matched: int = 0
    review: int = 0
    errors: list[dict] = field(default_factory=list)


def _header_map(header: list) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, raw in enumerate(header):
        name = str(raw or "").strip().lower()
        for key, aliases in COLUMN_ALIASES.items():
            if name in aliases and key not in mapping:
                mapping[key] = idx
    if "title" not in mapping and "article" not in mapping:
        raise DomainError("В файле нет колонок «Наименование» или «Артикул»")
    if "price" not in mapping:
        raise DomainError("В файле нет колонки «Цена»")
    return mapping


def read_table(filename: str, content: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    elif name.endswith((".csv", ".txt")):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1251")
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=";,\t")
        rows = list(csv.reader(io.StringIO(text), dialect))
    else:
        raise DomainError("Поддерживаются файлы .xlsx и .csv")
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        raise DomainError("Файл пуст")
    mapping = _header_map(rows[0])
    return [{k: (r[i] if i < len(r) else None) for k, i in mapping.items()} for r in rows[1:]]


def _int(value, field_name: str) -> int:
    if value is None or str(value).strip() == "":
        raise ValueError(f"не заполнено поле «{field_name}»")
    s = str(value).replace(" ", "").replace(" ", "").replace(",", ".")
    n = float(s)
    if n < 0:
        raise ValueError(f"«{field_name}» не может быть отрицательным")
    return round(n)


def _str(value) -> str | None:
    s = str(value).strip() if value is not None else ""
    if s.endswith(".0") and s[:-2].isdigit():  # Excel превращает артикулы-числа во float
        s = s[:-2]
    return s or None


def apply_rows(
    db: Session, shop: Shop, rows: list[dict], *, stock_mode: str, movement_kind: str, user_id: int | None
) -> ImportReport:
    report = ImportReport()
    warehouses = {w.name.strip().lower(): w for w in db.scalars(select(Warehouse).where(Warehouse.shop_id == shop.id))}
    default_wh = inv.default_warehouse_id(db, shop)

    for n, row in enumerate(rows, start=2):  # 1-я строка файла — заголовок
        report.rows += 1
        sp = db.begin_nested()
        try:
            sku, article, title = _str(row.get("sku")), _str(row.get("article")), _str(row.get("title"))
            price = _int(row.get("price"), "Цена")
            raw_qty = row.get("qty")
            qty = _int(raw_qty, "Количество") if raw_qty is not None and str(raw_qty).strip() != "" else None
            wh_name = (_str(row.get("warehouse")) or "").lower()
            if wh_name and wh_name not in warehouses:
                raise ValueError(f"склад «{row.get('warehouse')}» не найден")
            wh_id = warehouses[wh_name].id if wh_name else default_wh

            offer = None
            if sku:
                offer = db.scalar(select(Offer).where(Offer.shop_id == shop.id, Offer.sku == sku))
            if offer is None and article:
                offer = db.scalar(
                    select(Offer).where(Offer.shop_id == shop.id, Offer.article_normalized == normalize_number(article))
                )
            if offer is None:
                if not title:
                    raise ValueError("для нового товара нужно наименование")
                offer = Offer(shop_id=shop.id, sku=sku or inv.next_sku(db, shop), title=title, price=price)
                db.add(offer)
            is_new = offer.id is None
            offer.title = title or offer.title
            offer.price = price
            offer.article = article or offer.article
            offer.manufacturer = _str(row.get("manufacturer")) or offer.manufacturer
            offer.stock_mode = stock_mode
            offer.is_active = True
            auto_match(db, offer)
            db.flush()

            if qty is not None:
                stock = inv.get_or_create_stock(db, offer, wh_id)
                if _str(row.get("bin")):
                    stock.bin_location = _str(row.get("bin"))
                if stock.on_hand != qty:
                    inv.set_count(db, stock, qty, user_id=user_id, kind=movement_kind, note="Загрузка из файла" if movement_kind == "import" else "Синхронизация")
            sp.commit()
            if is_new:
                report.created += 1
            else:
                report.updated += 1
            if offer.match_status == "matched":
                report.matched += 1
            else:
                report.review += 1
        except (ValueError, DomainError, IntegrityError) as exc:
            sp.rollback()
            msg = "конфликт данных (дубликат кода?)" if isinstance(exc, IntegrityError) else getattr(exc, "message", str(exc))
            report.errors.append({"row": n, "error": msg})
    return report
